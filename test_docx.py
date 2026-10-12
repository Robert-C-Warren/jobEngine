import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from pathlib import Path

def test_docx():
    doc = docx.Document()
    
    # Set standard 0.6 inch margins for professional 1-2 page resume
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
        
    # Name
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_after = Pt(2)
    run_name = p_name.add_run("Robert Warren")
    run_name.font.name = "Calibri"
    run_name.font.size = Pt(22)
    run_name.font.bold = True
    run_name.font.color.rgb = RGBColor(18, 52, 86) # Deep Navy
    
    # Contact
    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_after = Pt(8)
    run_contact = p_contact.add_run("Dayton, OH  •  (626) 872-4171  •  robertcwarren1@gmail.com  •  linkedin.com/in/robert-c-warren")
    run_contact.font.name = "Calibri"
    run_contact.font.size = Pt(10)
    run_contact.font.color.rgb = RGBColor(80, 80, 80)
    
    # Save
    out_path = Path("test_resume.docx")
    doc.save(out_path)
    print("DOCX created successfully:", out_path.exists())

if __name__ == "__main__":
    test_docx()
